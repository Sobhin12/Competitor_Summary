"""Formula / Inputs trace columns: rendering, and that compute_derived_metrics
returns a trace for every derived cell whose inputs reproduce its value."""
import pytest

from competitor_analysis import config as cfg

if cfg.FY is None:
    cfg.set_period("FY25-26", "Q3")

from competitor_analysis.extraction import data_engine as de
from competitor_analysis.extraction.trace import Inp, Trace, fmt_num


def test_fmt_num_indian_grouping_and_missing():
    assert fmt_num(429500.0) == "4,29,500.00"
    assert fmt_num(1234567) == "12,34,567"
    assert fmt_num(0.0786) == "0.0786"
    assert fmt_num(None) == "n/a"
    assert fmt_num(-1500.5) == "-1,500.50"


def test_trace_render_has_cur_prior_and_note():
    t = Trace("A / B", [Inp("A", 10.0, "NL-4", "Cr")], [Inp("A", None, "NL-4", "Cr")], note="fallback used")
    formula, inputs = t.render()
    assert formula == "A / B"
    assert inputs == "CUR: A [NL-4] 10.00 Cr | PRIOR: A [NL-4] n/a Cr | NOTE: fallback used"


def test_make_trace_cur_only_omits_prior():
    t = de.make_trace("x", [de._inp("A", "NL-1", "L", 5, 7)], cur_only=True)
    assert t.prior == [] and "PRIOR" not in t.render()[1]


def test_write_trace_and_clear_roundtrip():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    de.write_trace(ws, 2, Trace("f", [Inp("A", 1.0)]))
    assert ws.cell(2, 11).value == "f" and "CUR: A 1.00" in ws.cell(2, 12).value
    de.write_trace(ws, 2, "PRIOR: last-year fact", append=True)
    assert ws.cell(2, 12).value.endswith("| PRIOR: last-year fact")
    assert ws.cell(2, 11).value == "f"


def test_sync_period_headers_writes_trace_headers():
    import openpyxl
    ws = openpyxl.Workbook().active
    de.sync_period_headers(ws)
    assert ws.cell(1, 11).value == "Formula" and ws.cell(1, 12).value == "Inputs"


def test_apply_metric_to_rows_writes_trace_only_with_value():
    import openpyxl
    ws = openpyxl.Workbook().active
    for c, v in enumerate(["Slide #", "Category", "Company", "Meric 1", "Metric 2", "s", "l", "CUR", "PRIOR", "Growth"], 1):
        ws.cell(1, c).value = v
    ws.append([24, "cat", "NBHI", "Opex. To GWP ratio", None])
    idx = de.build_row_index(ws)
    t = Trace("a / b", [Inp("a", 2.0)])
    de.apply_metric_to_rows(ws, idx, 24, "NBHI", "Opex. To GWP ratio", None, 0.5, 0.25, False, [], trace=t)
    assert ws.cell(2, 11).value == "a / b"


def test_clear_period_values_clears_trace_columns():
    import openpyxl
    ws = openpyxl.Workbook().active
    for c, v in enumerate(["Slide #", "Category", "Company", "Meric 1", "Metric 2", "s", "l", "CUR", "PRIOR", "Growth"], 1):
        ws.cell(1, c).value = v
    ws.append([24, "cat", "NBHI", "m", None, None, None, 1.0, 2.0, 0.5, "f", "inputs"])
    de.clear_period_values(ws)
    assert ws.cell(2, 11).value is None and ws.cell(2, 12).value is None


def _derive_with_traces(income, raw=None):
    from competitor_analysis.extraction import gemini, schemas
    specs = gemini.master_metric_specs()
    regrouped = schemas.regroup_by_form(raw or {}, specs)
    kind_by_key = {m["key"]: m["kind"] for m in specs}
    traces = {}
    derived = de.compute_derived_metrics("ACME", regrouped, kind_by_key, income, opening_onroll=None, traces=traces)
    return derived, traces


def _found(cur, prior):
    return {"fy26_q3": cur, "fy25_q3": prior, "found": True, "source_form": "NL-7", "page_number": 1,
            "evidence": "e", "notes": None}


def test_traces_do_not_change_values_and_cover_derived_cells():
    income = {"gwp": (1000.0, 800.0), "opex_alone": (100.0, 90.0), "nwp": (900.0, 700.0),
              "ep": (850.0, 650.0), "net_commission": (50.0, 40.0), "net_incurred_claims": (500.0, 380.0),
              "pbt": (30.0, 20.0), "pat": (25.0, 15.0)}
    raw = {"manpower_cost": _found(50000.0, 40000.0), "capital": _found(1000.0, 900.0),
           "bs_reserves_surplus": _found(2000.0, 1800.0), "bs_debit_balance_pl": _found(300.0, 400.0)}
    from competitor_analysis.extraction import gemini, schemas
    specs = gemini.master_metric_specs()
    kind = {m["key"]: m["kind"] for m in specs}
    without = de.compute_derived_metrics("ACME", schemas.regroup_by_form(raw, specs), kind, income, opening_onroll=None)
    derived, traces = _derive_with_traces(income, raw)
    assert derived == without  # traces are additive
    # every value-bearing derived cell has a trace
    missing = [k for k, (c, p) in derived.items() if (c is not None or p is not None) and k not in traces]
    assert missing == []


def test_trace_inputs_reproduce_the_value():
    income = {"gwp": (1000.0, 800.0), "opex_alone": (100.0, 90.0), "nwp": (900.0, 700.0),
              "ep": (850.0, 650.0), "net_commission": (50.0, 40.0), "net_incurred_claims": (500.0, 380.0)}
    raw = {"manpower_cost": _found(50000.0, 40000.0), "capital": _found(1000.0, 900.0),
           "bs_reserves_surplus": _found(2000.0, 1800.0), "bs_debit_balance_pl": _found(300.0, 400.0)}
    derived, traces = _derive_with_traces(income, raw)

    # Manpower to GWP = Manpower used / GWP, straight from the trace's inputs
    t = traces[(24, "Manpower to GWP ratio", None)]
    used = next(i for i in t.cur if i.label == "Manpower used").value
    gwp = next(i for i in t.cur if i.label == "Gross Written Premium").value
    assert round(used / gwp, 4) == derived[(24, "Manpower to GWP ratio", None)][0]

    # Net Worth = (capital + reserves - debit) x 100
    t = traces[(26, "Net Worth", None)]
    v = {i.label: i.value for i in t.cur}
    assert round((v["Share Capital"] + v["Reserves & Surplus"] - v["Debit balance in P&L"]) * 100, 2) == \
        derived[(26, "Net Worth", None)][0]

    # Expense ratio = (net commission + opex) / NWP
    t = traces[(22, "Expense Ratio", None)]
    v = {i.label: i.value for i in t.prior}
    assert round((v["Net Commission"] + v["Operating Expenses"]) / v["Net Written Premium"], 4) == \
        derived[(22, "Expense Ratio", None)][1]
